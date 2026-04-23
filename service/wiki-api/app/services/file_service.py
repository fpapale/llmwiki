import os
from .path_service import PathService

class FileService:
    def __init__(self, path_service: PathService):
        self.path_service = path_service

    def read_file(self, file_path: str) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
            
        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.pdf':
            import pypdf
            text = ""
            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                for page in reader.pages:
                    text += (page.extract_text() or "") + "\n"
            return text
        elif ext == '.docx':
            import docx
            doc = docx.Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs])
        elif ext == '.xlsx':
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = ""
            for sheet in wb.worksheets:
                text += f"--- Sheet: {sheet.title} ---\n"
                for row in sheet.iter_rows(values_only=True):
                    # Filter out None values and join row data
                    row_text = "\t".join([str(cell) for cell in row if cell is not None])
                    if row_text.strip():
                        text += row_text + "\n"
            return text
        else:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()

    def write_file(self, file_path: str, content: str) -> None:
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)

    def append_file(self, file_path: str, content: str) -> None:
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "a", encoding="utf-8") as f:
            f.write(content)
            
    def list_files(self, dir_path: str, extension: str = None) -> list[str]:
        if not os.path.exists(dir_path):
            return []
        files = []
        for root, _, filenames in os.walk(dir_path):
            for filename in filenames:
                if extension and not filename.endswith(extension):
                    continue
                files.append(os.path.join(root, filename))
        return files
